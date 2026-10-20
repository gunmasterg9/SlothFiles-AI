import email
from email.message import EmailMessage
from pathlib import Path
import docx
import fitz
import openpyxl
import pytest

from backend.app.services.document_extractor import document_extractor


def test_extract_plain_text(tmp_path: Path):
    """Test text extraction from plain text and markdown files."""
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Hello SlothPilot AI. Testing extractor.", encoding="utf-8")

    res = document_extractor.extract_text(txt_file)
    assert res.has_text is True
    assert "SlothPilot AI" in res.text_content
    assert res.ocr_required is False


def test_extract_docx(tmp_path: Path):
    """Test text extraction from DOCX file."""
    docx_file = tmp_path / "document.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph of confidential agreement.")
    doc.add_paragraph("Second paragraph with invoice details.")
    doc.save(str(docx_file))

    res = document_extractor.extract_text(docx_file)
    assert res.has_text is True
    assert "confidential agreement" in res.text_content
    assert "invoice details" in res.text_content


def test_extract_xlsx(tmp_path: Path):
    """Test text extraction from XLSX workbook."""
    xlsx_file = tmp_path / "sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(["Item", "Quantity", "Amount"])
    ws.append(["Widgets", 50, 1250.00])
    wb.save(str(xlsx_file))

    res = document_extractor.extract_text(xlsx_file)
    assert res.has_text is True
    assert "Widgets" in res.text_content
    assert "1250" in res.text_content


def test_extract_eml(tmp_path: Path):
    """Test text extraction from standard EML email message."""
    eml_file = tmp_path / "message.eml"
    msg = EmailMessage()
    msg["Subject"] = "Project Milestone Update"
    msg["From"] = "alice@example.com"
    msg["To"] = "bob@example.com"
    msg.set_content("Please find attached the signed contract. Deadline is Friday.")

    with open(eml_file, "wb") as f:
        f.write(msg.as_bytes())

    res = document_extractor.extract_text(eml_file)
    assert res.has_text is True
    assert "Project Milestone Update" in res.text_content
    assert "alice@example.com" in res.text_content
    assert "Deadline is Friday" in res.text_content


def test_extract_pdf_and_scanned_detection(tmp_path: Path):
    """Test text extraction from PDF and detection of scanned documents."""
    # 1. Text PDF
    pdf_file = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "SlothPilot Bill of Lading BL-9901")
    doc.save(str(pdf_file))
    doc.close()

    res = document_extractor.extract_text(pdf_file)
    assert res.has_text is True
    assert "BL-9901" in res.text_content
    assert res.ocr_required is False

    # 2. Empty/scanned PDF without text layer
    scanned_file = tmp_path / "scanned.pdf"
    doc2 = fitz.open()
    doc2.new_page()  # Blank page without text
    doc2.save(str(scanned_file))
    doc2.close()

    res2 = document_extractor.extract_text(scanned_file)
    assert res2.has_text is False
    assert res2.ocr_required is True
    assert "OCR" in (res2.warning or "")
