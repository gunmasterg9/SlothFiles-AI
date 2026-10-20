import email
from email import policy
import logging
from pathlib import Path
from pydantic import BaseModel

import fitz  # PyMuPDF
import docx
import openpyxl

from backend.app.config import settings

logger = logging.getLogger("slothpilot.extractor")


class ExtractionResult(BaseModel):
    file_path: str
    file_name: str
    extension: str
    size_bytes: int
    text_content: str
    page_or_sheet_count: int = 1
    has_text: bool = True
    ocr_required: bool = False
    warning: str | None = None


class DocumentExtractor:
    """Extracts text content and metadata from supported document formats."""

    SUPPORTED_EXTENSIONS = {
        ".txt", ".md", ".json", ".csv", ".log",
        ".pdf", ".docx", ".xlsx", ".eml"
    }

    @classmethod
    def can_extract(cls, path: Path) -> bool:
        return path.suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def extract_text(cls, path: Path) -> ExtractionResult:
        resolved = path.resolve()
        ext = resolved.suffix.lower()
        file_size = resolved.stat().st_size

        if file_size > settings.max_file_size_bytes:
            return ExtractionResult(
                file_path=str(resolved),
                file_name=resolved.name,
                extension=ext,
                size_bytes=file_size,
                text_content="",
                has_text=False,
                warning=f"File exceeds maximum allowed size ({settings.max_file_size_bytes} bytes)."
            )

        try:
            if ext in {".txt", ".md", ".json", ".csv", ".log"}:
                return cls._extract_plain_text(resolved, ext, file_size)
            elif ext == ".pdf":
                return cls._extract_pdf(resolved, ext, file_size)
            elif ext == ".docx":
                return cls._extract_docx(resolved, ext, file_size)
            elif ext == ".xlsx":
                return cls._extract_xlsx(resolved, ext, file_size)
            elif ext == ".eml":
                return cls._extract_eml(resolved, ext, file_size)
            else:
                return ExtractionResult(
                    file_path=str(resolved),
                    file_name=resolved.name,
                    extension=ext,
                    size_bytes=file_size,
                    text_content="",
                    has_text=False,
                    warning="Unsupported file format."
                )
        except Exception as exc:
            logger.warning("Extraction failed for %s: %s", resolved.name, exc)
            return ExtractionResult(
                file_path=str(resolved),
                file_name=resolved.name,
                extension=ext,
                size_bytes=file_size,
                text_content="",
                has_text=False,
                warning=f"Extraction failed: {type(exc).__name__} ({str(exc)})"
            )

    @staticmethod
    def _extract_plain_text(path: Path, ext: str, size: int) -> ExtractionResult:
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="latin-1", errors="replace")
        return ExtractionResult(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            size_bytes=size,
            text_content=text[:20000],
            has_text=bool(text.strip())
        )

    @staticmethod
    def _extract_pdf(path: Path, ext: str, size: int) -> ExtractionResult:
        doc = fitz.open(str(path))
        pages_text = []
        page_count = len(doc)
        total_chars = 0

        for i in range(min(page_count, 50)):  # Bounded to first 50 pages
            page_text = doc[i].get_text()
            total_chars += len(page_text.strip())
            pages_text.append(page_text)
        doc.close()

        full_text = "\n\n".join(pages_text).strip()
        # If PDF has pages but 0 extracted characters, it is likely a scanned image
        ocr_required = (page_count > 0 and total_chars == 0)

        return ExtractionResult(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            size_bytes=size,
            page_or_sheet_count=page_count,
            text_content=full_text[:20000],
            has_text=bool(full_text),
            ocr_required=ocr_required,
            warning="Scanned PDF detected. Local OCR engine is not configured." if ocr_required else None
        )

    @staticmethod
    def _extract_docx(path: Path, ext: str, size: int) -> ExtractionResult:
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs).strip()
        return ExtractionResult(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            size_bytes=size,
            page_or_sheet_count=1,
            text_content=full_text[:20000],
            has_text=bool(full_text)
        )

    @staticmethod
    def _extract_xlsx(path: Path, ext: str, size: int) -> ExtractionResult:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        sheet_texts = []
        sheet_count = len(wb.sheetnames)

        for sheet_name in wb.sheetnames[:5]:  # Bounded to 5 sheets
            sheet = wb[sheet_name]
            rows_text = []
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                row_count += 1
                if row_count > 100:  # Bounded to 100 rows per sheet
                    break
                row_cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if row_cells:
                    rows_text.append(" | ".join(row_cells))
            if rows_text:
                sheet_texts.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows_text))
        wb.close()

        full_text = "\n\n".join(sheet_texts).strip()
        return ExtractionResult(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            size_bytes=size,
            page_or_sheet_count=sheet_count,
            text_content=full_text[:20000],
            has_text=bool(full_text)
        )

    @staticmethod
    def _extract_eml(path: Path, ext: str, size: int) -> ExtractionResult:
        with open(path, "rb") as f:
            msg = email.message_from_binary_file(f, policy=policy.default)

        subject = msg.get("Subject", "")
        sender = msg.get("From", "")
        recipient = msg.get("To", "")
        date_str = msg.get("Date", "")

        body_parts = []
        body = msg.get_body(preferencelist=("plain", "html"))
        if body:
            content = body.get_content()
            body_parts.append(content)

        full_text = f"Subject: {subject}\nFrom: {sender}\nTo: {recipient}\nDate: {date_str}\n\n" + "\n\n".join(body_parts)
        return ExtractionResult(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            size_bytes=size,
            page_or_sheet_count=1,
            text_content=full_text[:20000],
            has_text=bool(full_text.strip())
        )


document_extractor = DocumentExtractor()
